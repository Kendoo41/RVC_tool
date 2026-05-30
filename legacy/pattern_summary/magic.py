#!/common/appl/python/python-3.7.7/bin/python3.7
"""
Pattern Extraction (v3) - Linux-run, Windows-readable hyperlinks
================================================================
Key changes from v2:
- Path mapping table to translate Linux abs paths to Windows UNC / drive letters.
- Optional 'relative path' mode (recommended if VIL files & summary file
  are kept in the same project tree).
- Backslash conversion to satisfy Windows Excel hyperlink syntax.
"""

import os
import re
import glob
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ----------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------
INPUT_DIR   = "./1_VIL"
OUTPUT_FILE = "./Pattern_Summary.xlsx"

# Hyperlink mode:
#   "relative" : link by relative path (best if VILs + Summary share a folder tree)
#   "mapped"   : translate Linux abs path -> Windows path via PATH_MAP below
#   "unc"      : prefix Linux abs path with a UNC server (PATH_MAP single entry)
HYPERLINK_MODE = "mapped"

# Linux -> Windows path mapping. EDIT THIS to match your IT setup.
# Order matters: longer / more specific prefixes should come first.
# Examples:
#   "/proj/hpc/u2b20ez"  -> "\\\\fileserver\\proj\\hpc\\u2b20ez"
#   "/home/khoidao"      -> "H:\\khoidao"
#   "/scratch"           -> "S:\\scratch"
# Linux -> Windows path mapping
PATH_MAP = [
    # ("linux_prefix", "windows_prefix")
    ("/shsv", r"\\rvc-vnas-01.rvc.renesas.com"),
]

# Priority ranking (higher = higher priority)
PRIORITY_RANK = {"S": 3, "A": 2, "B": 1, "": 0, None: 0}

SKIP_SHEETS_EXACT    = {"summary", "info", "cover"}
SKIP_SHEETS_CONTAINS = ["change history", "testcase list", "rtl testcase",
                         "gate testcase", "history"]

# ----------------------------------------------------------------------
# Path translation: Linux -> Windows
# ----------------------------------------------------------------------
def to_windows_path(linux_abs_path, output_dir_abs):
    """
    Translate a Linux absolute path into a string Windows Excel can open.
    Returns a path string suitable for hyperlink target (without the '#sheet!cell').
    """
    linux_abs_path = os.path.abspath(linux_abs_path)

    if HYPERLINK_MODE == "relative":
        # Relative path: works on any OS as long as both files keep relative layout.
        rel = os.path.relpath(linux_abs_path, start=output_dir_abs)
        # Excel/Windows: use backslashes
        return rel.replace("/", "\\")

    if HYPERLINK_MODE in ("mapped", "unc"):
        # Sort PATH_MAP by descending prefix length to match the most specific first
        for lin_pref, win_pref in sorted(PATH_MAP, key=lambda x: -len(x[0])):
            # Normalize trailing slashes
            lp = lin_pref.rstrip("/")
            if linux_abs_path == lp or linux_abs_path.startswith(lp + "/"):
                tail = linux_abs_path[len(lp):]            # starts with "/..."
                tail_win = tail.replace("/", "\\")          # convert separators
                # Make sure we don't get double-backslashes at the join
                wp = win_pref.rstrip("\\")
                return wp + tail_win

        # No mapping matched -> warn and fall back to raw path with backslashes
        print(f"  [WARN] No PATH_MAP entry matched: {linux_abs_path}")
        print(f"         Hyperlink may not open on Windows. "
              f"Add a mapping entry covering this prefix.")
        return linux_abs_path.replace("/", "\\")

    # Unknown mode -> raw fallback
    return linux_abs_path.replace("/", "\\")


def build_hyperlink_target(linux_file, sheet_name, row, output_dir_abs):
    """
    Build a full Excel hyperlink target string:
        <windows_path>#'<sheet>'!A<row>
    """
    win_path = to_windows_path(linux_file, output_dir_abs)
    # Sheet names with spaces/special chars need single quotes
    return f"{win_path}#'{sheet_name}'!A{row}"


# ----------------------------------------------------------------------
# (everything below is identical to v2 except the hyperlink construction)
# ----------------------------------------------------------------------
def cell_text(cell_data_only, cell_with_formula):
    v1 = cell_data_only.value if cell_data_only is not None else None
    v2 = cell_with_formula.value if cell_with_formula is not None else None
    out = []
    for v in (v1, v2):
        if v is None:
            continue
        s = str(v)
        if s.startswith("="):
            s = s[1:]
        out.append(s)
    return " ".join(out)


def row_text(ws_data, ws_form, row_idx, max_col):
    parts = []
    for c in range(1, max_col + 1):
        parts.append(cell_text(ws_data.cell(row=row_idx, column=c),
                               ws_form.cell(row=row_idx, column=c)))
    return parts


def is_skippable_sheet(name):
    n = name.strip().lower()
    if n in SKIP_SHEETS_EXACT:
        return True
    return any(k in n for k in SKIP_SHEETS_CONTAINS)


def find_header_rows(ws_data, ws_form, max_scan=120):
    max_col = min(ws_data.max_column or 1, 80)
    max_row = min(ws_data.max_row or 1, max_scan)
    for r in range(1, max_row + 1):
        row_vals = row_text(ws_data, ws_form, r, max_col)
        joined = " | ".join(v.lower() for v in row_vals)
        if ("main item" in joined and "name of test pattern" in joined
            and "confirmation" in joined):
            header_map = {}
            for c_idx, val in enumerate(row_vals, start=1):
                key = val.strip()
                if key and key.lower() != "skip":
                    header_map[key] = c_idx
            sub_r = r + 1
            if sub_r <= (ws_data.max_row or 1):
                sub_vals = row_text(ws_data, ws_form, sub_r, max_col)
                for c_idx, val in enumerate(sub_vals, start=1):
                    key = val.strip()
                    if key and key.lower() != "skip" and key not in header_map:
                        header_map[key] = c_idx
            return r, sub_r, header_map
    return None, None, None


def get_col(header_map, *candidates):
    for cand in candidates:
        cand_l = cand.lower()
        for k, v in header_map.items():
            if cand_l == k.lower():
                return v
        for k, v in header_map.items():
            if cand_l in k.lower():
                return v
    return None


PATTERN_NAME_RE = re.compile(r"^[A-Za-z0-9_][A-Za-z0-9_.\-]*$")

def split_patterns(cell_value):
    if not cell_value:
        return []
    text = str(cell_value).replace("_x000D_", "")\
                          .replace("\\n", "\n").replace("\r", "\n")
    parts = re.split(r"[,\n;]+", text)
    out = []
    for p in parts:
        p = p.strip()
        if not p or p.lower() in {"-", "skip", "—", "tbd"}:
            continue
        if not PATTERN_NAME_RE.match(p):
            continue
        out.append(p)
    return out


def priority_max(p1, p2):
    return p1 if PRIORITY_RANK.get(p1, 0) >= PRIORITY_RANK.get(p2, 0) else p2


def scan_vil_files(input_dir):
    patterns_info = defaultdict(lambda: {"priority": "", "usages": []})
    file_paths = sorted(glob.glob(os.path.join(input_dir, "*.xls*")))
    for fpath in file_paths:
        print(f"[INFO] Scanning: {os.path.basename(fpath)}")
        try:
            wb_data = load_workbook(fpath, data_only=True)
            wb_form = load_workbook(fpath, data_only=False)
        except Exception as e:
            print(f"  [WARN] Cannot open: {e}")
            continue
        for sheet_name in wb_data.sheetnames:
            if is_skippable_sheet(sheet_name):
                continue
            if sheet_name not in wb_form.sheetnames:
                continue
            ws_data = wb_data[sheet_name]
            ws_form = wb_form[sheet_name]
            main_r, sub_r, header_map = find_header_rows(ws_data, ws_form)
            if main_r is None:
                continue
            col_pattern  = get_col(header_map, "Name of test pattern", "test pattern")
            col_confirm  = get_col(header_map, "Confirmation")
            col_main     = get_col(header_map, "Main item")
            col_middle   = get_col(header_map, "Middle item")
            col_detail   = get_col(header_map, "Detailed item")
            col_priority = get_col(header_map, "priority")
            if not col_pattern:
                continue
            data_start = (sub_r or main_r) + 1
            print(f"  [OK] '{sheet_name}' header={main_r}, "
                  f"data_start={data_start}, prio_col={col_priority}")
            n_items = 0
            for r in range(data_start, (ws_data.max_row or 0) + 1):
                pattern_cell = ws_data.cell(row=r, column=col_pattern).value
                if pattern_cell is None:
                    pattern_cell = ws_form.cell(row=r, column=col_pattern).value
                pats = split_patterns(pattern_cell)
                if not pats:
                    continue
                parts = []
                for c in (col_main, col_middle, col_detail):
                    if c:
                        v = ws_data.cell(row=r, column=c).value
                        if v in (None, ""):
                            v = ws_form.cell(row=r, column=c).value
                        if v not in (None, "", "skip"):
                            parts.append(str(v).strip())
                if col_confirm:
                    v = ws_data.cell(row=r, column=col_confirm).value
                    if v in (None, ""):
                        v = ws_form.cell(row=r, column=col_confirm).value
                    if v:
                        txt = str(v).strip().replace("\n", " ")\
                                            .replace("_x000D_", "")
                        if len(txt) > 200:
                            txt = txt[:200] + "..."
                        parts.append(txt)
                checkpoint = " | ".join(parts) if parts else "(no description)"
                row_prio = ""
                if col_priority:
                    pv = ws_data.cell(row=r, column=col_priority).value
                    if pv in (None, ""):
                        pv = ws_form.cell(row=r, column=col_priority).value
                    if pv:
                        row_prio = str(pv).strip().upper()
                        if row_prio not in PRIORITY_RANK:
                            row_prio = ""
                for pat in pats:
                    patterns_info[pat]["priority"] = priority_max(
                        patterns_info[pat]["priority"], row_prio)
                    patterns_info[pat]["usages"].append({
                        "file": os.path.abspath(fpath),
                        "file_name": os.path.basename(fpath),
                        "sheet": sheet_name,
                        "row": r,
                        "checkpoint": checkpoint,
                    })
                n_items += 1
            print(f"     -> {n_items} item rows")
        wb_data.close()
        wb_form.close()
    return patterns_info, file_paths


def build_output(patterns_info, output_path):
    output_abs   = os.path.abspath(output_path)
    output_dir   = os.path.dirname(output_abs)

    wb = Workbook()
    thin = Side(border_style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    link_font = Font(color="0563C1", underline="single")
    prio_fill = {
        "S": PatternFill("solid", fgColor="FFC7CE"),
        "A": PatternFill("solid", fgColor="FFEB9C"),
        "B": PatternFill("solid", fgColor="C6EFCE"),
        "":  PatternFill("solid", fgColor="EEEEEE"),
    }
    sorted_patterns = sorted(patterns_info.keys(), key=lambda x: x.lower())

    # ----- Summary sheet (internal links, no path conversion needed) -----
    ws = wb.active
    ws.title = "Summary"
    ws.append(["No.", "Pattern Name", "Priority", "Usage Count", "Detail"])
    for c in range(1, 6):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for idx, pat in enumerate(sorted_patterns, start=1):
        info = patterns_info[pat]
        prio = info["priority"] or ""
        ws.cell(row=idx+1, column=1, value=idx).border = border
        ws.cell(row=idx+1, column=2, value=pat).border = border
        pc = ws.cell(row=idx+1, column=3, value=prio)
        pc.fill = prio_fill.get(prio, prio_fill[""])
        pc.alignment = Alignment(horizontal="center"); pc.border = border
        ws.cell(row=idx+1, column=4, value=len(info["usages"])).border = border
        #lc = ws.cell(row=idx+1, column=5, value=f"Open No. {idx}")
        # lc.hyperlink = f"#'No. {idx}'!A1"     # internal link, no OS issue
        # lc.font = link_font; lc.border = border
        # New (internal link via HYPERLINK formula):
        ws.cell(row=idx+1, column=5,
                value=f'=HYPERLINK("#\'No. {idx}\'!A1","Open No. {idx}")')

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 16
    ws.freeze_panes = "A2"

    # ----- Detail sheets (EXTERNAL hyperlinks - need Windows path) -----
    for idx, pat in enumerate(sorted_patterns, start=1):
        info = patterns_info[pat]
        sheet_name = f"No. {idx}"[:31]
        ds = wb.create_sheet(title=sheet_name)
        ds["A1"] = "Pattern Name:"; ds["B1"] = pat
        ds["A2"] = "Priority:";     ds["B2"] = info["priority"] or "(none)"
        ds["A3"] = "Total usages:"; ds["B3"] = len(info["usages"])
        # New:
        ds["A4"] = '=HYPERLINK("#\'Summary\'!A1","Back to Summary")'
        ds["A4"].font = link_font
        for r in range(1, 4):
            ds.cell(row=r, column=1).font = Font(bold=True)

        headers = ["#", "Check Point (Item)", "Source VIL", "Sheet",
                   "Row", "Open Source"]
        for c_idx, h in enumerate(headers, start=1):
            c = ds.cell(row=6, column=c_idx, value=h)
            c.font = header_font; c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border

        for u_idx, u in enumerate(info["usages"], start=1):
            row = 6 + u_idx
            ds.cell(row=row, column=1, value=u_idx).border = border
            ds.cell(row=row, column=2, value=u["checkpoint"]).border = border
            ds.cell(row=row, column=2).alignment = Alignment(
                wrap_text=True, vertical="top")
            ds.cell(row=row, column=3, value=u["file_name"]).border = border
            ds.cell(row=row, column=4, value=u["sheet"]).border = border
            ds.cell(row=row, column=5, value=u["row"]).border = border

            lc = ds.cell(row=row, column=6, value="Open")
            # >>> Windows-friendly hyperlink target <<<
            # Build target path (Windows UNC)
            target = build_hyperlink_target(u["file"], u["sheet"], u["row"], output_dir)

            # Use HYPERLINK formula instead of hyperlink object - more reliable for UNC paths
            # Escape any double-quote in target (UNC paths shouldn't have any, but just in case)
            target_escaped = target.replace('"', '""')
            lc = ds.cell(row=row, column=6,
                         value=f'=HYPERLINK("{target_escaped}","Open")')
            lc.font = link_font
            lc.border = border

        ds.column_dimensions["A"].width = 5
        ds.column_dimensions["B"].width = 80
        ds.column_dimensions["C"].width = 45
        ds.column_dimensions["D"].width = 35
        ds.column_dimensions["E"].width = 8
        ds.column_dimensions["F"].width = 12
        ds.freeze_panes = "A7"

    wb.save(output_path)
    print(f"[DONE] Output -> {output_path}")
    print(f"[INFO] Hyperlink mode: {HYPERLINK_MODE}")
    if HYPERLINK_MODE == "mapped":
        print("[INFO] Path mappings used:")
        for lp, wp in PATH_MAP:
            print(f"         {lp}  ->  {wp}")


def main():
    if not os.path.isdir(INPUT_DIR):
        os.makedirs(INPUT_DIR, exist_ok=True)
        print(f"[INFO] Created {INPUT_DIR}. Put VIL files there and re-run.")
        return
    patterns_info, files = scan_vil_files(INPUT_DIR)
    if not patterns_info:
        print("[WARN] No patterns extracted.")
        return
    print(f"[INFO] Files       : {len(files)}")
    print(f"[INFO] Patterns    : {len(patterns_info)}")
    build_output(patterns_info, OUTPUT_FILE)


if __name__ == "__main__":
    main()