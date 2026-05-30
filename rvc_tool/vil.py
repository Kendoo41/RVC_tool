"""
vil.py - Scan Verification Item List (VIL) workbooks.
=====================================================

Ported from the original ``magic.py``. Instead of writing an Excel summary
directly, this module returns a plain data structure that the rest of the tool
joins with the master report and pattern sources::

    {
      "<pattern_name>": {
          "priority": "S" | "A" | "B" | "",
          "usages": [
              {"file", "file_name", "sheet", "row", "checkpoint",
               "main", "middle", "detailed", "confirmation"},
              ...
          ],
      },
      ...
    }

The header-detection heuristic is unchanged: a row containing "main item",
"name of test pattern" and "confirmation" marks the column layout, and data
rows start two lines below.
"""
from __future__ import annotations

import glob
import os
import re
from collections import defaultdict
from typing import Dict, List, Optional

from openpyxl import load_workbook

PRIORITY_RANK = {"S": 3, "A": 2, "B": 1, "": 0, None: 0}

SKIP_SHEETS_EXACT = {"summary", "info", "cover"}
SKIP_SHEETS_CONTAINS = [
    "change history", "testcase list", "rtl testcase",
    "gate testcase", "history",
]

PATTERN_NAME_RE = re.compile(r"^[A-Za-z0-9_][A-Za-z0-9_.\-]*$")


def _cell_text(cell_data_only, cell_with_formula) -> str:
    v1 = cell_data_only.value if cell_data_only is not None else None
    v2 = cell_with_formula.value if cell_with_formula is not None else None
    out = []
    for v in (v1, v2):
        if v is None:
            continue
        s = str(v)
        if s.startswith("="):
            s = s[1:]
        out.append(s)
    return " ".join(out)


def _row_text(ws_data, ws_form, row_idx, max_col) -> List[str]:
    return [
        _cell_text(ws_data.cell(row=row_idx, column=c), ws_form.cell(row=row_idx, column=c))
        for c in range(1, max_col + 1)
    ]


def _is_skippable_sheet(name: str) -> bool:
    n = name.strip().lower()
    if n in SKIP_SHEETS_EXACT:
        return True
    return any(k in n for k in SKIP_SHEETS_CONTAINS)


def _find_header_rows(ws_data, ws_form, max_scan=120):
    max_col = min(ws_data.max_column or 1, 80)
    max_row = min(ws_data.max_row or 1, max_scan)
    for r in range(1, max_row + 1):
        row_vals = _row_text(ws_data, ws_form, r, max_col)
        joined = " | ".join(v.lower() for v in row_vals)
        if ("main item" in joined and "name of test pattern" in joined
                and "confirmation" in joined):
            header_map: Dict[str, int] = {}
            for c_idx, val in enumerate(row_vals, start=1):
                key = val.strip()
                if key and key.lower() != "skip":
                    header_map[key] = c_idx
            sub_r = r + 1
            if sub_r <= (ws_data.max_row or 1):
                sub_vals = _row_text(ws_data, ws_form, sub_r, max_col)
                for c_idx, val in enumerate(sub_vals, start=1):
                    key = val.strip()
                    if key and key.lower() != "skip" and key not in header_map:
                        header_map[key] = c_idx
            return r, sub_r, header_map
    return None, None, None


def _get_col(header_map: Dict[str, int], *candidates) -> Optional[int]:
    for cand in candidates:
        cand_l = cand.lower()
        for k, v in header_map.items():
            if cand_l == k.lower():
                return v
        for k, v in header_map.items():
            if cand_l in k.lower():
                return v
    return None


def split_patterns(cell_value) -> List[str]:
    """Split a 'Name of test pattern' cell into individual pattern names."""
    if not cell_value:
        return []
    text = (str(cell_value).replace("_x000D_", "")
            .replace("\\n", "\n").replace("\r", "\n"))
    parts = re.split(r"[,\n;]+", text)
    out = []
    for p in parts:
        p = p.strip()
        if not p or p.lower() in {"-", "skip", "—", "tbd"}:
            continue
        if not PATTERN_NAME_RE.match(p):
            continue
        out.append(p)
    return out


def _priority_max(p1, p2):
    return p1 if PRIORITY_RANK.get(p1, 0) >= PRIORITY_RANK.get(p2, 0) else p2


def scan_vils(input_dir=None, files=None, progress=None) -> Dict[str, dict]:
    """
    Scan every ``*.xls*`` file in ``input_dir`` (or the explicit ``files`` list)
    and return ``{pattern_name: {"priority", "usages": [...]}}``.

    ``progress`` is an optional ``callable(str)`` for log output.
    """
    def log(msg):
        if progress:
            progress(msg)

    patterns_info: Dict[str, dict] = defaultdict(lambda: {"priority": "", "usages": []})

    if files is None:
        files = sorted(glob.glob(os.path.join(input_dir or ".", "*.xls*")))
    # Ignore Excel lock files (~$...).
    files = [f for f in files if not os.path.basename(f).startswith("~$")]

    for fpath in files:
        log("[VIL] scanning {}".format(os.path.basename(fpath)))
        try:
            wb_data = load_workbook(fpath, data_only=True)
            wb_form = load_workbook(fpath, data_only=False)
        except Exception as e:  # pragma: no cover - corrupt file guard
            log("  [WARN] cannot open: {}".format(e))
            continue

        for sheet_name in wb_data.sheetnames:
            if _is_skippable_sheet(sheet_name):
                continue
            if sheet_name not in wb_form.sheetnames:
                continue
            ws_data = wb_data[sheet_name]
            ws_form = wb_form[sheet_name]
            main_r, sub_r, header_map = _find_header_rows(ws_data, ws_form)
            if main_r is None:
                continue

            col_pattern = _get_col(header_map, "Name of test pattern", "test pattern")
            col_confirm = _get_col(header_map, "Confirmation")
            col_main = _get_col(header_map, "Main item")
            col_middle = _get_col(header_map, "Middle item")
            col_detail = _get_col(header_map, "Detailed item")
            col_priority = _get_col(header_map, "priority")
            if not col_pattern:
                continue

            data_start = (sub_r or main_r) + 1
            for r in range(data_start, (ws_data.max_row or 0) + 1):
                pattern_cell = ws_data.cell(row=r, column=col_pattern).value
                if pattern_cell is None:
                    pattern_cell = ws_form.cell(row=r, column=col_pattern).value
                pats = split_patterns(pattern_cell)
                if not pats:
                    continue

                def field_at(col):
                    if not col:
                        return ""
                    v = ws_data.cell(row=r, column=col).value
                    if v in (None, ""):
                        v = ws_form.cell(row=r, column=col).value
                    if v in (None, "", "skip"):
                        return ""
                    return str(v).strip().replace("_x000D_", "")

                main = field_at(col_main)
                middle = field_at(col_middle)
                detailed = field_at(col_detail)
                confirmation = field_at(col_confirm).replace("\n", " ")
                if len(confirmation) > 400:
                    confirmation = confirmation[:400] + "..."

                parts = [p for p in (main, middle, detailed, confirmation) if p]
                checkpoint = " | ".join(parts) if parts else "(no description)"

                row_prio = ""
                if col_priority:
                    pv = field_at(col_priority)
                    if pv:
                        row_prio = pv.strip().upper()
                        if row_prio not in PRIORITY_RANK:
                            row_prio = ""

                for pat in pats:
                    info = patterns_info[pat]
                    info["priority"] = _priority_max(info["priority"], row_prio)
                    info["usages"].append({
                        "file": os.path.abspath(fpath),
                        "file_name": os.path.basename(fpath),
                        "sheet": sheet_name,
                        "row": r,
                        "checkpoint": checkpoint,
                        "main": main,
                        "middle": middle,
                        "detailed": detailed,
                        "confirmation": confirmation,
                    })

        wb_data.close()
        wb_form.close()

    log("[VIL] {} pattern(s) extracted".format(len(patterns_info)))
    return dict(patterns_info)
