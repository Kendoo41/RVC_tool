"""
Tests for the RVC tool. Runnable with either::

    python -m pytest tests/
    python tests/test_rvc.py        # falls back to a tiny built-in runner
"""
from __future__ import annotations

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from rvc_tool import listfile as lf
from rvc_tool import master_report as mr
from rvc_tool import model as modelmod
from rvc_tool import vil as vilmod


# ---------------------------------------------------------------------------
# listfile
# ---------------------------------------------------------------------------
SAMPLE_LIST = """\
# set default option
--no-clean
--fsimopts -define FAST_SIM
-fcc
#========== PE0 ==========
HBUS_TOP/sub/cpuss_addr_spi0   --fsimopts -vlog_t foo.v
HBUS_TOP/sub/cpuss_addr_spi1
#HBUS_TOP/sub/user_disabled_pattern --fsimopts bar
INT_TOP/INT_reg_init --fsimopts -simvopt "+CODE_NAME=INT_reg_init "
"""


def test_classify():
    entries = [lf.classify_line(l) for l in SAMPLE_LIST.splitlines()]
    kinds = [e.kind for e in entries]
    assert kinds.count(lf.OPTION) == 3          # --no-clean, --fsimopts, -fcc
    assert kinds.count(lf.COMMENT) == 2         # "# set default option", "#==== PE0 ===="
    assert kinds.count(lf.PATTERN) == 4         # 3 active + 1 user-commented
    # bare name extraction strips the hierarchy - no module arg needed
    names = lf.pattern_names(entries)
    assert "cpuss_addr_spi0" in names
    assert "INT_reg_init" in names
    assert "user_disabled_pattern" in names     # commented pattern still discovered


def test_filter_roundtrip():
    with tempfile.NamedTemporaryFile("w", suffix=".list", delete=False) as fh:
        fh.write(SAMPLE_LIST)
        path = fh.name
    try:
        original = open(path).read()
        status = {
            "cpuss_addr_spi0": "PASS",
            "cpuss_addr_spi1": "FAIL",
            "INT_reg_init": "NA",
        }
        res = lf.update_listfile(path, status, comment_statuses=("PASS", "FAIL", "NA"))
        assert res.commented.get("PASS") == 1
        assert res.commented.get("FAIL") == 1
        assert res.commented.get("NA") == 1
        text = open(path).read()
        assert "#PSS HBUS_TOP/sub/cpuss_addr_spi0" in text
        assert "#NG  HBUS_TOP/sub/cpuss_addr_spi1" in text
        # user-commented line must be left exactly as-is
        assert "#HBUS_TOP/sub/user_disabled_pattern --fsimopts bar" in text

        # restore => byte-identical to the original
        lf.update_listfile(path, {}, restore=True)
        assert open(path).read() == original
    finally:
        os.unlink(path)


def test_filter_keep():
    with tempfile.NamedTemporaryFile("w", suffix=".list", delete=False) as fh:
        fh.write(SAMPLE_LIST)
        path = fh.name
    try:
        status = {"cpuss_addr_spi0": "PASS", "cpuss_addr_spi1": "FAIL", "INT_reg_init": "PASS"}
        # keep only FAIL active => PASS get commented, FAIL stays
        comment = [s for s in ("PASS", "FAIL", "NA", "MISSING") if s != "FAIL"]
        lf.update_listfile(path, status, comment_statuses=comment)
        text = open(path).read()
        assert "#PSS HBUS_TOP/sub/cpuss_addr_spi0" in text
        assert text.count("\nHBUS_TOP/sub/cpuss_addr_spi1") == 1  # still active
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# master_report
# ---------------------------------------------------------------------------
def test_report_parsing():
    d = tempfile.mkdtemp()
    older = os.path.join(d, "old.rpt")
    newer = os.path.join(d, "new.rpt")
    open(older, "w").write(
        "==== Execution Results ====\n"
        "  [OK]   (12:00:01) HBUS_TOP/sub/patA --fsimopts foo\n"
        "  [NG]   (12:00:02) HBUS_TOP/sub/patB --fsimopts foo\n"
    )
    open(newer, "w").write(
        "  [N/A]  (12:01:01) HBUS_TOP/sub/patC --fsimopts foo\n"
        "  [OK]   (12:01:02) HBUS_TOP/sub/patB --fsimopts foo\n"   # newer overrides patB
    )
    os.utime(older, (1000, 1000))
    os.utime(newer, (2000, 2000))

    reports = mr.parse_reports([d])
    sm = mr.status_map(reports)
    assert sm["patA"] == "PASS"
    assert sm["patB"] == "PASS"     # newest .rpt wins (was FAIL in older)
    assert sm["patC"] == "NA"


def test_report_name_extraction():
    assert mr._test_name("HBUS_TOP/sub/foo_bar --fsimopts x") == "foo_bar"
    assert mr._test_name("foo_bar") == "foo_bar"
    assert mr._test_name("a/b/c/name -opt") == "name"


# ---------------------------------------------------------------------------
# vil
# ---------------------------------------------------------------------------
def _make_vil(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.cell(row=3, column=3, value="Main item")
    ws.cell(row=3, column=5, value="Middle item")
    ws.cell(row=3, column=7, value="Detailed item")
    ws.cell(row=3, column=9, value="Confirmation")
    ws.cell(row=3, column=10, value="Name of test pattern")
    ws.cell(row=3, column=11, value="Priority")
    ws.cell(row=5, column=3, value="Reg")
    ws.cell(row=5, column=9, value="Check reg RW")
    ws.cell(row=5, column=10, value="pat_one, pat_two")   # two patterns in one cell
    ws.cell(row=5, column=11, value="S")
    wb.save(path)


def test_vil_scan():
    d = tempfile.mkdtemp()
    path = os.path.join(d, "demo_verification_item_list.xlsx")
    _make_vil(path)
    info = vilmod.scan_vils(files=[path])
    assert "pat_one" in info and "pat_two" in info
    assert info["pat_one"]["priority"] == "S"
    cp = info["pat_one"]["usages"][0]
    assert cp["main"] == "Reg"
    assert "Check reg RW" in cp["confirmation"]


# ---------------------------------------------------------------------------
# model join
# ---------------------------------------------------------------------------
def test_model_join():
    d = tempfile.mkdtemp()
    list_path = os.path.join(d, "S_item.list")
    open(list_path, "w").write("pat_one\npat_two\npat_three\n")
    vil_info = {
        "pat_one": {"priority": "S", "usages": [{"main": "X", "middle": "", "detailed": "",
                                                  "confirmation": "do x", "checkpoint": "X | do x",
                                                  "file_name": "v.xlsx", "sheet": "s", "row": 5}]},
        # planned in the VIL but in NO list and NO report -> the coverage gap
        "pat_vil_only": {"priority": "A", "usages": [{"main": "Y", "middle": "", "detailed": "",
                                                       "confirmation": "do y", "checkpoint": "Y | do y",
                                                       "file_name": "v.xlsx", "sheet": "s", "row": 9}]},
    }
    reports = {
        "pat_one": {"status": "PASS", "rpt": "a.rpt", "rpt_dir": "v1"},
        "pat_two": {"status": "FAIL", "rpt": "a.rpt", "rpt_dir": "v1"},
        # pat_three absent -> MISSING
        "pat_unlisted": {"status": "PASS", "rpt": "a.rpt", "rpt_dir": "v1"},
    }
    ds = modelmod.build_dataset([list_path], vil_info=vil_info, reports=reports)
    # LIST viewpoint totals must NOT count the VIL-only item (it's not listed
    # and has no report line) -> same numbers as before this feature.
    t = ds.totals()
    assert t["PASS"] == 2 and t["FAIL"] == 1 and t["MISSING"] == 1
    assert t["TOTAL"] == 4  # pat_one, pat_two, pat_three, pat_unlisted (NOT pat_vil_only)
    p1 = ds.patterns["pat_one"]
    assert p1.priority == "S" and p1.status == "PASS" and p1.checkpoints
    assert p1.in_vil is True and p1.in_list is True
    assert any(p.name == "pat_unlisted" for p in ds.unlisted)

    # VIL-only item exists, is flagged in_vil but not in_list, and is MISSING.
    vo = ds.patterns["pat_vil_only"]
    assert vo.in_vil is True and vo.in_list is False and vo.status == "MISSING"

    # VIL viewpoint totals: 2 planned items, 1 listed, 1 not-listed (the gap).
    vt = ds.vil_totals()
    assert vt["TOTAL"] == 2 and vt["LISTED"] == 1 and vt["NOT_LISTED"] == 1
    # vil_groups: S bucket has pat_one, A bucket has pat_vil_only.
    buckets = {g.priority: [p.name for p in g.patterns] for g in ds.vil_groups}
    assert buckets.get("S") == ["pat_one"] and buckets.get("A") == ["pat_vil_only"]

    # priority item lists
    pl = modelmod.priority_item_lists(vil_info)
    assert pl["S"] == ["pat_one"]


# ---------------------------------------------------------------------------
# tiny standalone runner
# ---------------------------------------------------------------------------
def _run_all():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    failed = 0
    for t in tests:
        try:
            t()
            print("  PASS  {}".format(t.__name__))
        except Exception as e:  # noqa
            failed += 1
            print("  FAIL  {}: {}".format(t.__name__, e))
    print("\n{}/{} passed".format(len(tests) - failed, len(tests)))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(_run_all())
