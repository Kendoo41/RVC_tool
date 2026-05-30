#!/usr/bin/env python3
"""
make_demo.py - Generate a small, self-contained demo fixture set.
=================================================================

Creates synthetic (non-proprietary) data under ``examples/demo/`` so the tool
can be exercised end-to-end without any real project files::

    examples/demo/
      vil/demo_HBUS_verification_item_list.xlsx   (VIL with the expected header)
      master_report/v001_001/run_a.rpt            ([OK]/[NG]/[N/A] result lines)
      master_report/v001_002/run_b.rpt
      testcase/HBUS_TOP/.../<pattern>/<files>     (asm/stimulus/assertion source)
      lists/S_item.list  A_item.list  B_item.list (bare pattern names)

Run:  python examples/make_demo.py
Then: python -m rvc_tool build -c examples/rvc.example.json
"""
from __future__ import annotations

import os
import random

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DEMO = os.path.join(HERE, "demo")

# A handful of made-up patterns grouped by priority.
PATTERNS = {
    "S": [
        ("hbus_qos_reg_rw", "Register", "QoS", "Read/Write all QoS registers", "Confirm RW of every QoS register bit"),
        ("hbus_qos_arbiter_priority", "Function", "Arbiter", "Priority ordering", "High-priority master wins arbitration"),
        ("hbus_safety_ecc_1bit", "Safety", "ECC", "1-bit error correction", "Single-bit error is corrected, no interrupt"),
    ],
    "A": [
        ("hbus_qos_bandwidth_limit", "Function", "Bandwidth", "Limiter cap", "Throughput stays under programmed cap"),
        ("hbus_err_response_slverr", "Error", "Response", "SLVERR path", "Illegal access returns SLVERR"),
        ("hbus_safety_ecc_2bit", "Safety", "ECC", "2-bit error detection", "Double-bit error raises the safety interrupt"),
        ("hbus_addr_map_flash", "Address", "Map", "Flash region", "Access lands in the flash address window"),
    ],
    "B": [
        ("hbus_reg_default_value", "Register", "Reset", "Default values", "All registers read their reset default"),
        ("hbus_clock_gating", "Function", "Clock", "Gating", "Clock gates when idle"),
        ("hbus_addr_map_sram", "Address", "Map", "SRAM region", "Access lands in the SRAM address window"),
        ("hbus_low_power_retention", "Power", "Retention", "Deep stop", "State retained across deep-stop"),
        ("hbus_dummy_unrun_pattern", "Function", "Misc", "Not executed yet", "Placeholder - no result in the report"),
    ],
}


def make_vil():
    out_dir = os.path.join(DEMO, "vil")
    os.makedirs(out_dir, exist_ok=True)
    wb = Workbook()
    # Skippable sheet (exercises the skip logic).
    cover = wb.active
    cover.title = "Summary"
    cover["B2"] = "Demo VIL - Summary (skipped by the scanner)"

    def add_sheet(title, rows):
        ws = wb.create_sheet(title)
        # Header at row 3, sub-header row 4, data from row 5 - matching the
        # real layout that vil.py auto-detects.
        ws.cell(row=3, column=3, value="Main item")
        ws.cell(row=3, column=5, value="Middle item")
        ws.cell(row=3, column=7, value="Detailed item")
        ws.cell(row=3, column=9, value="Confirmation")
        ws.cell(row=3, column=10, value="Name of test pattern")
        ws.cell(row=3, column=11, value="Priority")
        ws.cell(row=4, column=9, value="contents")
        ws.cell(row=4, column=10, value="assembler")
        r = 5
        for (name, main, middle, detail, conf, prio) in rows:
            ws.cell(row=r, column=3, value=main)
            ws.cell(row=r, column=5, value=middle)
            ws.cell(row=r, column=7, value=detail)
            ws.cell(row=r, column=9, value=conf)
            ws.cell(row=r, column=10, value=name)
            ws.cell(row=r, column=11, value=prio)
            r += 1

    qos_rows = []
    safety_rows = []
    for prio, items in PATTERNS.items():
        for (name, main, middle, detail, conf) in items:
            target = safety_rows if "safety" in name or "err" in name else qos_rows
            target.append((name, main, middle, detail, conf, prio))
    add_sheet("hbus_qos", qos_rows)
    add_sheet("hbus_safety", safety_rows)

    path = os.path.join(out_dir, "demo_HBUS_verification_item_list.xlsx")
    wb.save(path)
    return path


def make_reports():
    random.seed(7)
    all_names = [n for items in PATTERNS.values() for (n, *_rest) in items]
    # Assign statuses; leave the "dummy_unrun" one out -> MISSING.
    status = {}
    for n in all_names:
        if n.endswith("unrun_pattern"):
            continue
        rnd = random.random()
        status[n] = "OK" if rnd < 0.6 else ("NG" if rnd < 0.8 else "N/A")

    d1 = os.path.join(DEMO, "master_report", "v001_001")
    d2 = os.path.join(DEMO, "master_report", "v001_002")
    os.makedirs(d1, exist_ok=True)
    os.makedirs(d2, exist_ok=True)

    def line(tag, name):
        return "  [{}]   (12:0{}:{:02d}) HBUS_TOP/sub/{}   --fsimopts -simvopt foo\n".format(
            tag, random.randint(0, 9), random.randint(0, 59), name)

    a, b = [], []
    for n, tag in status.items():
        (a if random.random() < 0.5 else b).append(line(tag, n))
    hdr = "==== Execution Results ====\n"
    open(os.path.join(d1, "run_a.rpt"), "w").write(hdr + "".join(a))
    open(os.path.join(d2, "run_b.rpt"), "w").write(hdr + "".join(b))
    return [d2, d1]  # newest-ish first


def make_patterns():
    base = os.path.join(DEMO, "testcase", "HBUS_TOP", "sub")
    made = 0
    for items in PATTERNS.values():
        for (name, *_rest) in items:
            if name.endswith("unrun_pattern"):
                continue  # no source for the not-run example
            d = os.path.join(base, name)
            os.makedirs(d, exist_ok=True)
            open(os.path.join(d, name + ".s"), "w").write(
                "// asm source for {0}\n_start_{0}:\n    li   a0, 0x4000_0000\n"
                "    sw   a0, 0(a1)      // program QoS register\n    ret\n".format(name))
            open(os.path.join(d, name + ".v"), "w").write(
                "// stimulus for {0}\nmodule {0}_stim;\n  initial begin\n"
                "    #10 force top.dut.qos_en = 1'b1;\n    #100 release top.dut.qos_en;\n"
                "  end\nendmodule\n".format(name))
            open(os.path.join(d, "checker.sv"), "w").write(
                "// assertion for {0}\nproperty p_{0};\n  @(posedge clk) req |-> ##[1:3] ack;\n"
                "endproperty\nassert property (p_{0});\n".format(name))
            made += 1
    return base, made


def make_lists():
    out_dir = os.path.join(DEMO, "lists")
    os.makedirs(out_dir, exist_ok=True)
    for prio, items in PATTERNS.items():
        path = os.path.join(out_dir, "{}_item.list".format(prio))
        with open(path, "w") as fh:
            fh.write("# {} priority demo list\n".format(prio))
            for (name, *_rest) in items:
                fh.write(name + "\n")
    return out_dir


def main():
    os.makedirs(DEMO, exist_ok=True)
    vil = make_vil()
    rpt_dirs = make_reports()
    pat_base, n = make_patterns()
    lists = make_lists()
    print("[demo] VIL        :", vil)
    print("[demo] reports    :", ", ".join(rpt_dirs))
    print("[demo] patterns   : {} folders under {}".format(n, pat_base))
    print("[demo] lists      :", lists)
    print("[demo] now run    : python -m rvc_tool build -c examples/rvc.example.json")


if __name__ == "__main__":
    main()
