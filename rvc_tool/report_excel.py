"""
report_excel.py - Build the AI-review workbook.
===============================================

Two views (plus a summary), as chosen:

  * **By Checkpoint** - one row per (pattern x checkpoint). This is the sheet an
    AI (or a human) uses to judge whether each VIL checkpoint is actually
    reflected in the pattern. It carries two empty columns,
    "Reflected? (AI)" and "Evidence / Notes (AI)", for the reviewer to fill.
  * **By Pattern** - one row per pattern: priority, status, lists, the joined
    checkpoints, source file names and (truncated) source content.
  * **Summary** - totals per status and per priority.
"""
from __future__ import annotations

from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import Dataset, Pattern

_THIN = Side(border_style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="305496")
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center")

_STATUS_FILL = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "NA": PatternFill("solid", fgColor="D9D9D9"),
    "MISSING": PatternFill("solid", fgColor="FFEB9C"),
}
_PRIO_FILL = {
    "S": PatternFill("solid", fgColor="F8CBAD"),
    "A": PatternFill("solid", fgColor="FFE699"),
    "B": PatternFill("solid", fgColor="DDEBF7"),
}
_STATUS_DISP = {"PASS": "PASS", "FAIL": "FAIL", "NA": "N/A", "MISSING": "-"}
_SOURCE_CAP = 30000  # cap concatenated source per pattern cell


def _header(ws, headers: List[str]):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.freeze_panes = "A2"


def _ordered_patterns(ds: Dataset) -> List[Pattern]:
    return sorted(ds.patterns.values(), key=lambda p: p.sort_key())


def _by_checkpoint_sheet(wb, ds: Dataset):
    ws = wb.active
    ws.title = "By Checkpoint"
    headers = ["Pattern", "Priority", "Status", "List(s)",
               "Main item", "Middle item", "Detailed item", "Confirmation",
               "VIL file", "Sheet", "Row", "Source files",
               "Reflected? (AI)", "Evidence / Notes (AI)"]
    _header(ws, headers)

    r = 2
    for p in _ordered_patterns(ds):
        src_files = ", ".join(f.get("file", "") for f in p.files)
        lists = ", ".join(p.lists)
        rows = p.checkpoints or [{}]
        for c in rows:
            ws.cell(row=r, column=1, value=p.name).border = _BORDER
            pc = ws.cell(row=r, column=2, value=p.priority or "")
            pc.alignment = _CENTER
            pc.border = _BORDER
            if p.priority in _PRIO_FILL:
                pc.fill = _PRIO_FILL[p.priority]
            sc = ws.cell(row=r, column=3, value=_STATUS_DISP.get(p.status, p.status))
            sc.alignment = _CENTER
            sc.border = _BORDER
            if p.status in _STATUS_FILL:
                sc.fill = _STATUS_FILL[p.status]
            ws.cell(row=r, column=4, value=lists).border = _BORDER
            ws.cell(row=r, column=5, value=c.get("main", "")).border = _BORDER
            ws.cell(row=r, column=6, value=c.get("middle", "")).border = _BORDER
            ws.cell(row=r, column=7, value=c.get("detailed", "")).border = _BORDER
            conf = ws.cell(row=r, column=8, value=c.get("confirmation", ""))
            conf.border = _BORDER
            conf.alignment = _WRAP
            ws.cell(row=r, column=9, value=c.get("file_name", "")).border = _BORDER
            ws.cell(row=r, column=10, value=c.get("sheet", "")).border = _BORDER
            ws.cell(row=r, column=11, value=c.get("row", "")).border = _BORDER
            ws.cell(row=r, column=12, value=src_files).border = _BORDER
            ws.cell(row=r, column=13, value="").border = _BORDER       # Reflected? (AI)
            ws.cell(row=r, column=14, value="").border = _BORDER       # Evidence (AI)
            r += 1

    widths = [34, 8, 8, 18, 22, 22, 28, 50, 30, 18, 6, 30, 14, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = "A1:N1"


def _by_pattern_sheet(wb, ds: Dataset):
    ws = wb.create_sheet("By Pattern")
    headers = ["Pattern", "Priority", "Status", "List(s)", "# Checkpoints",
               "Checkpoints (joined)", "Source files", "Sim options",
               "Source content (truncated)", "AI review"]
    _header(ws, headers)

    r = 2
    for p in _ordered_patterns(ds):
        cps = " || ".join(
            " | ".join(x for x in (c.get("main", ""), c.get("middle", ""),
                                   c.get("detailed", ""), c.get("confirmation", "")) if x)
            for c in p.checkpoints
        )
        src_files = ", ".join(f.get("file", "") for f in p.files)
        blob = ""
        for f in p.files:
            blob += "===== {} =====\n{}\n\n".format(f.get("file", ""), f.get("content", ""))
            if len(blob) > _SOURCE_CAP:
                blob = blob[:_SOURCE_CAP] + "\n... [truncated]"
                break

        ws.cell(row=r, column=1, value=p.name).border = _BORDER
        pc = ws.cell(row=r, column=2, value=p.priority or "")
        pc.alignment = _CENTER
        pc.border = _BORDER
        if p.priority in _PRIO_FILL:
            pc.fill = _PRIO_FILL[p.priority]
        sc = ws.cell(row=r, column=3, value=_STATUS_DISP.get(p.status, p.status))
        sc.alignment = _CENTER
        sc.border = _BORDER
        if p.status in _STATUS_FILL:
            sc.fill = _STATUS_FILL[p.status]
        ws.cell(row=r, column=4, value=", ".join(p.lists)).border = _BORDER
        ws.cell(row=r, column=5, value=len(p.checkpoints)).border = _BORDER
        cpc = ws.cell(row=r, column=6, value=cps)
        cpc.border = _BORDER
        cpc.alignment = _WRAP
        ws.cell(row=r, column=7, value=src_files).border = _BORDER
        oc = ws.cell(row=r, column=8, value=p.options)
        oc.border = _BORDER
        oc.alignment = _WRAP
        bc = ws.cell(row=r, column=9, value=blob)
        bc.border = _BORDER
        bc.alignment = _WRAP
        ws.cell(row=r, column=10, value="").border = _BORDER
        r += 1

    widths = [34, 8, 8, 18, 12, 60, 30, 40, 70, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = "A1:J1"


def _summary_sheet(wb, ds: Dataset):
    ws = wb.create_sheet("Summary", 0)
    totals = ds.totals()
    ws["A1"] = "RVC Verification Summary"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Status"
    ws["B3"] = "Count"
    for cell in (ws["A3"], ws["B3"]):
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
    rows = [("Total", totals["TOTAL"]), ("Pass", totals["PASS"]),
            ("Fail", totals["FAIL"]), ("N/A", totals["NA"]), ("Not run", totals["MISSING"])]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=k).border = _BORDER
        ws.cell(row=i, column=2, value=v).border = _BORDER

    # priority breakdown
    prio = {"S": 0, "A": 0, "B": 0, "(none)": 0}
    for p in ds.patterns.values():
        prio[p.priority if p.priority in ("S", "A", "B") else "(none)"] += 1
    ws["D3"] = "Priority"
    ws["E3"] = "Count"
    for cell in (ws["D3"], ws["E3"]):
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
    for i, k in enumerate(("S", "A", "B", "(none)"), start=4):
        ws.cell(row=i, column=4, value=k).border = _BORDER
        ws.cell(row=i, column=5, value=prio[k]).border = _BORDER

    # per-list breakdown
    ws["A10"] = "Per list"
    ws["A10"].font = Font(bold=True)
    hdr = ["List", "Total", "Pass", "Fail", "N/A", "Not run"]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=11, column=c, value=h)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
    r = 12
    for g in ds.groups:
        c = g.counts()
        vals = [g.label, len(g.patterns), c["PASS"], c["FAIL"], c["NA"], c["MISSING"]]
        for ci, v in enumerate(vals, start=1):
            ws.cell(row=r, column=ci, value=v).border = _BORDER
        r += 1

    for col, w in (("A", 22), ("B", 10), ("C", 10), ("D", 12), ("E", 10), ("F", 10)):
        ws.column_dimensions[col].width = w


def generate(ds: Dataset, out_xlsx: str) -> str:
    """Write the workbook and return its path."""
    wb = Workbook()
    _by_checkpoint_sheet(wb, ds)   # uses wb.active (becomes sheet 2 after summary insert)
    _by_pattern_sheet(wb, ds)
    _summary_sheet(wb, ds)         # inserted at index 0
    wb.save(out_xlsx)
    return out_xlsx
